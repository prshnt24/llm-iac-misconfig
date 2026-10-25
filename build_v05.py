"""Build Prashant_Singh_Draft_v05.docx from v04: new chapters 5 and 6, revised methodology, appendices and references."""
import copy, re, json, pathlib, sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from thesis_data import *
from thesis_text import T, rem
import thesis_text_rev  # applies the post-review revisions to T
import llm_evaluation as ev
import judge_findings as jf
import remediation_judge as rj
import validity_judge as vj

SRC, OUT = ROOT / "Prashant_Singh_Draft_v04.docx", ROOT / "Prashant_Singh_Draft_v05.docx"
doc = Document(SRC)
body = doc.element.body


def paras():
    return doc.paragraphs


def find(prefix, start=0):
    for p in paras():
        if p.text.startswith(prefix):
            return p
    raise KeyError(prefix)


def settext(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(text)
    return p


def sub_in(p, old, new):
    done = False
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            done = True
    if not done and old in p.text:
        settext(p, p.text.replace(old, new))


def P(el):
    return Paragraph(el, doc._body)


# ---------- templates (deep copies taken before any edit) ----------
tpl = {
    "body": copy.deepcopy(find("Checkov reported 221 failed checks")._p),
    "h2": copy.deepcopy(find("5.3 RQ1: Overall Detection Accuracy")._p),
    "h1": copy.deepcopy(find("CHAPTER 6: CONCLUSION")._p),
    "cap": copy.deepcopy(find("Table 5.1. Static analysis findings")._p),
    "pb": copy.deepcopy([e for e in body if e.tag.endswith("}p") and list(e.iter(qn("w:br"))) and not "".join(t.text or "" for t in e.iter(qn("w:t")))][-1]),
    "ref": copy.deepcopy(find("Anthropic. (2026, February 17)")._p),
    "h2_app": copy.deepcopy(find("B.1 Zero-shot prompt")._p),
}


class Cur:
    def __init__(self, el):
        self.el = el

    def put(self, el):
        self.el.addnext(el)
        self.el = el
        return el

    def para(self, text, key="body"):
        el = copy.deepcopy(tpl[key])
        self.put(el)
        p = P(el)
        settext(p, text)
        return p

    def h2(self, t):
        return self.para(t, "h2")

    def h1(self, t):
        return self.para(t, "h1")

    def pagebreak(self):
        self.put(copy.deepcopy(tpl["pb"]))

    def caption(self, t):
        return self.para(t, "cap")

    def code(self, text):
        p = self.para(text, "body")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs:
            r.font.name = "Courier New"
            r.font.size = Pt(8.5)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        return p

    def figure(self, path, caption, width=6.0):
        el = copy.deepcopy(tpl["body"])
        self.put(el)
        p = P(el)
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(6)
        p.add_run().add_picture(str(path), width=Inches(width))
        r = p.add_run()
        r.add_break()
        c = p.add_run(caption)
        c.italic = True
        c.font.name = "Times New Roman"
        c.font.size = Pt(11)
        return p

    def table(self, headers, rows, widths=None, size=9):
        t = doc.add_table(rows=1, cols=len(headers))
        for j, h in enumerate(headers):
            t.rows[0].cells[j].text = h
        for row in rows:
            cells = t.add_row().cells
            for j, val in enumerate(row):
                cells[j].text = str(val)
        total = 9360
        widths = widths or [1] * len(headers)
        ws = [int(total * w / sum(widths)) for w in widths]
        for i, row in enumerate(t.rows):
            for j, cell in enumerate(row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                for e in tcPr.findall(qn("w:tcW")):
                    tcPr.remove(e)
                w = OxmlElement("w:tcW")
                w.set(qn("w:w"), str(ws[j]))
                w.set(qn("w:type"), "dxa")
                tcPr.insert(0, w)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(0)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(size)
                        r.bold = (i == 0)
        self.put(t._tbl)
        sp = copy.deepcopy(tpl["body"])
        self.put(sp)
        settext(P(sp), "")
        P(sp).paragraph_format.line_spacing = 1.0
        return t


def apa_borders(t):
    tblPr = t._tbl.tblPr
    for e in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(e)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "bottom"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)
    for c in t.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr()
        for e in tcPr.findall(qn("w:tcBorders")):
            tcPr.remove(e)
        tb = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        tb.append(e)
        tcPr.append(tb)


def delete_between(start_el, end_el):
    """Delete start_el and following siblings up to (not including) end_el."""
    el = start_el
    while el is not None and el is not end_el:
        nxt = el.getnext()
        el.getparent().remove(el)
        el = nxt


# ---------- 0. global text hygiene ----------
for p in paras():
    for old, new in [("chapter?s", "chapter's")]:
        if old in p.text:
            sub_in(p, old, new)
    if "—" in p.text or "�" in p.text:
        t = p.text
        sep = ": " if t.startswith(("Figure", "Table")) else ", "
        t = re.sub(r"\s*[—�]\s*", sep, t)
        settext(p, t)
for tb in doc.tables:
    for row in tb.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if "—" in p.text or "�" in p.text:
                    for r in p.runs:
                        r.text = re.sub(r"\s*[—�]\s*", ", ", r.text)

# citation corrections
for p in paras():
    if "Rahman et al., 2023" in p.text:
        sub_in(p, "Rahman et al., 2023", "Rahman et al., 2019")
    if "Verdet et al., 2024" in p.text:
        sub_in(p, "Verdet et al., 2024", "Verdet et al., 2025")
for key, yr in [("Checkov is an open-source static analysis", "2020a"), ("The literature establishes", "2020a"), ("TerraGoat is a Terraform repository", "2020b")]:
    p = find(key)
    sub_in(p, "Bridgecrew, 2020", f"Bridgecrew, {yr}")

# ---------- 1. front matter and Chapter 1 ----------
settext(find("This dissertation investigates whether"), T["abstract"])
settext(find("The empirical scope is Terraform HCL"), T["scope"])
settext(find("Chapter 2 reviews IaC security"), T["structure"])

# ---------- 2. Chapter 2 ----------
p_iac = find("Infrastructure as Code is the practice of defining")
settext(p_iac, T["iac_survey"])
e_ = copy.deepcopy(p_iac._p); p_iac._p.addnext(e_); settext(P(e_), T["lit_iac_research"])
settext(find("LLMs are neural networks trained"), T["model_para"])
gap = find("No published study has conducted")
settext(gap, T["gap_hedge"])
cur = Cur(find("Vo et al. (2025) applied LLMs directly")._p)
cur.put(copy.deepcopy(gap._p)); settext(P(cur.el), T["lit_iac_eval"])
cur.put(copy.deepcopy(gap._p)); settext(P(cur.el), T["sec_llm_sec"][0])
h26 = find("2.6 Research Gap and Questions")
head_tpl, body_tpl = copy.deepcopy(h26._p), copy.deepcopy(gap._p)
settext(h26, "2.7 Research Gap and Questions")
c2 = Cur(h26._p.getprevious())
e = copy.deepcopy(head_tpl); c2.put(e); settext(P(e), T["sec_eval"][0])
e = copy.deepcopy(body_tpl); c2.put(e); settext(P(e), T["sec_eval"][1])
e = copy.deepcopy(body_tpl); c2.put(e); settext(P(e), T["lit_validity"])
sub_in(find("The literature establishes that Terraform"), "have not been evaluated specifically for Terraform HCL misconfiguration detection with quantitative metrics (Munshi et al., 2025; Vo et al., 2025).",
       "have not been evaluated specifically for Terraform HCL misconfiguration detection with quantitative metrics (Munshi et al., 2025; Vo et al., 2025). Work on judging LLM output supports the use of blinded judge models with human audit (Zheng et al., 2023).")

pg_ = find("This dissertation addresses that gap")
settext(pg_, T["rq_ref"])
for _ in range(3):
    nx = pg_._p.getnext()
    if nx is not None and P(nx).text.startswith("RQ"):
        nx.getparent().remove(nx)
e_ = copy.deepcopy(body_tpl); pg_._p.addprevious(e_); settext(P(e_), T["lit_synthesis"])

p_tool = find("Both tools have the same fundamental limitation")
e_ = copy.deepcopy(p_tool._p); p_tool._p.addnext(e_); settext(P(e_), T["lit_tools"])
p_ac = find("The closest related work is the ACSE-Eval")
e_ = copy.deepcopy(p_ac._p); p_ac._p.addnext(e_); settext(P(e_), T["lit_acse"])

# ---------- 3. Chapter 3 ----------
settext(find("This research uses a quantitative empirical design"), T["design_sentence"])
settext(find("A corpus of real-world Terraform repositories"), T["gh_note"])
settext(find("Approximately 30 to 50 community-contributed"), T["reg_note"])
settext(find("Ground truth is established using a two-tool consensus"), T["ground_truth"])
settext(find("The three models evaluated are GPT-4o"), T["models"])
settext(find("Two prompt engineering strategies are evaluated"), T["prompts"])
m = find("Each LLM response is parsed")
settext(m, T["matching"][0])
cm = Cur(m._p)
for t_ in T["matching"][1:]:
    e = copy.deepcopy(m._p); cm.put(e); settext(P(e), t_)
e = copy.deepcopy(m._p); cm.put(e); settext(P(e), T["llm_end"])
for hh in ("3.4.1 Model selection", "3.4.2 Prompt engineering approach", "3.4.3 Response parsing"):
    ph = find(hh)
    ph._p.getparent().remove(ph._p)
h34 = find("3.4 LLM Evaluation Design")
e = copy.deepcopy(m._p); h34._p.addnext(e); settext(P(e), T["llm_intro"])
settext(find("Precision is calculated as TP"), T["metrics"])
th = find("Four threats are acknowledged")
settext(th, T["threats"][0])
e = copy.deepcopy(th._p); th._p.addnext(e); settext(P(e), T["threats"][1])
da = find("All datasets, raw LLM responses")
settext(da, T["avail"])
hd = copy.deepcopy(find("3.7 Data Availability")._p)
c3 = Cur(da._p)
c3.put(hd); settext(P(hd), T["ai_decl"][0])
e = copy.deepcopy(da._p); c3.put(e); settext(P(e), T["ai_decl"][1])
t31 = doc.tables[0]
for row in t31.rows[1:]:
    if row.cells[0].text.startswith(("GitHub", "Terraform Registry")):
        row.cells[2].text = "30 .tf files (30 repositories)" if row.cells[0].text.startswith("GitHub") else "16 .tf files (16 modules)"
        for r in row.cells[2].paragraphs[0].runs:
            r.font.name = "Times New Roman"; r.font.size = Pt(12)
        if row.cells[0].text.startswith("Terraform Registry"):
            row.cells[1].text = "terraform-aws-modules organisation (GitHub source of Registry modules)"
            for r in row.cells[1].paragraphs[0].runs:
                r.font.name = "Times New Roman"; r.font.size = Pt(12)

sub_in(find("Three datasets are used."), "Three datasets are used.", "Three sources were planned. TerraGoat is the primary benchmark, and the GitHub and Registry files form one smaller generalisation set.")
lead31 = find("Figure 3.1 below provides a graphical overview")
settext(lead31, T["fig31_lead"])
figp = next(p for p in paras() if "Figure 3.1. Evaluation pipeline" in p.text)
for r_ in list(figp.runs):
    r_._r.getparent().remove(r_._r)
figp.alignment = WD_ALIGN_PARAGRAPH.CENTER
figp.paragraph_format.line_spacing = 1.0
figp.add_run().add_picture(str(ROOT / "terragoat/figure_3_1_methodology_pipeline_v2.png"), width=Inches(6.4))
figp.add_run().add_break()
cr_ = figp.add_run(T["fig31_caption"]); cr_.italic = True; cr_.font.name = "Times New Roman"; cr_.font.size = Pt(11)
anchor_ = figp._p
for t_ in T["fig31"]:
    e_ = copy.deepcopy(lead31._p); anchor_.addnext(e_); settext(P(e_), t_); anchor_ = e_

# ---------- 4. Chapter 4 ----------
sub_in(find("Chapter 4: Preliminary Results"), "Preliminary Results", "Static Analysis Results")
sub_in(find("This chapter presents findings from the first phase"),
       "These results form the preliminary ground truth baseline against which LLM detection performance will be compared. Trivy consensus scanning is the next step and will be completed before the LLM evaluation phase begins.",
       T["ch4_overview_tail"])
settext(find("4.8 Next Steps"), T["ch4_next_head"])
settext(find("The immediate next steps are"), T["ch4_next"])

for pref, new in T["captions"].items():
    if pref.startswith("Figure 4"):
        settext(find(pref), new)

# ---------- 5. Chapter 5 ----------
settext(find("This chapter reports the completed static analysis"), T["c5_intro"])
settext(find("The 75 candidate consensus groups are concentrated"), T["c5_cat_para"])
# rebuild Table 5.2 with the categorisation used for scoring
t52 = doc.tables[5]
c52 = Cur(t52._tbl.getprevious())
new52 = c52.table(["Category", "Groups", "Percentage"], [[c["category"], c["consensus_groups"], c["percentage"]] for c in CATSUM], widths=[3, 1, 1], size=12)
t52._tbl.getparent().remove(t52._tbl)

fig51 = find("Figure 5.1. Tool findings")
nxt = fig51._p.getnext()
if nxt is not None and list(nxt.iter(qn("w:drawing"))):
    fig51._p.addprevious(nxt)  # image above its caption, as for the other figures
cur = Cur(fig51._p)
settext(fig51, T["captions"]["Figure 5.1."])
cur.figure(ROOT / "terragoat/figure_5_2_consensus_categories.png", T["cap52"], 5.5)
anchor_after_fig = cur.el
ch5_start = find("5.3 RQ1: Overall Detection Accuracy")._p
refs_head = find("References")._p
delete_between(ch5_start, refs_head)
cur = Cur(anchor_after_fig)

TAB = "Table"
cur.h2(T["rq1_h"])
cur.para(T["rq1"][0])
cur.caption("Table 5.4. Detection performance against the 75 consensus findings")
rows = [[NAME[r.split("_")[0]], STRAT[r.split("_")[1]], S[r]["findings"], S[r]["strict_hits"], f2(S[r]["precision_strict"]), f2(S[r]["recall"]),
         f2(S[r]["f1_strict"]), f2(S[r]["precision_lenient"]), f2(S[r]["f1_lenient"])] for r in RUNS]
cur.table(["Model", "Prompt", "Findings", "Strict hits", "Precision (strict)", "Recall", "F1 (strict)", "Precision (lenient)", "F1 (lenient)"], rows, [2.2, 1.3, 1, 1, 1.2, 1, 1, 1.2, 1])
cur.figure(ROOT / "terragoat/figure_5_3_model_scores.png", T["cap53"], 6.0)
cur.para(T["rq1"][1])
cur.para(T["rq1"][2])
cur.para(T["rq1"][4])
cur.para(T["rq1_cov"])
cur.caption("Table 5.5. Coverage of the findings reported by Checkov and Trivy")
rows = [[NAME[r.split("_")[0]], STRAT[r.split("_")[1]], CV[r]["consensus_covered"], CV[r]["checkov_only_covered"], CV[r]["trivy_only_covered"],
         CV[r]["all_tool_rows_covered"], pc(CV[r]["of_223"])] for r in RUNS]
cur.table(["Model", "Prompt", f"Consensus (of {N_CONS})", f"Checkov only (of {N_CHK_ONLY})", f"Trivy only (of {N_TRV_ONLY})", f"All (of {N_ALL})", "Share"], rows, [2.2, 1.3, 1.3, 1.4, 1.3, 1.2, 1])
cur.para(T["rq1"][3])
cur.caption("Table 5.6. Exploratory validity check of findings that matched no tool finding")
rows = [[NAME[r.split("_")[0]], STRAT[r.split("_")[1]], S[r]["findings"], S[r]["lenient_hits"], VAL[r]["unmatched"], VAL[r]["valid"], VAL[r]["debatable"],
         VAL[r]["invalid"], f2(VAL[r]["adjusted_precision_valid"])] for r in RUNS]
cur.table(["Model", "Prompt", "Findings", "Matched a tool finding", "Unmatched", "Valid", "Arguable", "Invalid", "Adjusted precision"], rows, [2.2, 1.3, 1, 1.3, 1.2, 0.9, 1, 0.9, 1.3])

cur.h2(T["rq2_h"])
cur.para(T["rq2"][0])
cur.caption("Table 5.7. Recall by misconfiguration category")
rows = [[f"{c} (n = {gt_n(c)})"] + [f2(cat(r, c, "recall")) for r in RUNS] for c in CATS]
cur.table(["Category"] + [label(r).replace(", ", "\n") for r in RUNS], rows, [3, 1.2, 1.2, 1.2, 1.2, 1.2, 1.2])
cur.figure(ROOT / "terragoat/figure_5_4_category_recall.png", T["cap54"], 6.0)
cur.para(T["rq2"][1])
cur.para(T["rq2"][2])

cur.h2(T["rq3_h"])
cur.para(T["rq3"][0])
cur.caption("Table 5.8. Remediation quality of correctly detected findings (0 to 2 rubric)")
rows = []
for r in RUNS:
    d = REM_DIST[r]
    rows.append([NAME[r.split("_")[0]], STRAT[r.split("_")[1]], REM[r]["scored_findings"], rem(r), pc(REM[r]["gpt54_share_2"]), pc(REM[r]["sonnet55_share_2"]), f"{d[0]} / {d[1]} / {d[2]}"])
cur.table(["Model", "Prompt", "Items scored", "Mean score (both judges)", "Score 2 (judge 1)", "Score 2 (judge 2)", "Scores 0 / 1 / 2 (judge 1)"], rows, [2.2, 1.3, 1, 1, 1.2, 1.2, 1.7])
cur.para(T["rq3"][1])
cur.para(T["rq3"][2])

cur.h2(T["gen_h"])
cur.para(T["gen"][0])
cur.caption("Table 5.9. Generalisation check on public repositories and registry modules (zero-shot)")
rows = [[NAME[r.split("_")[0]], CS[r]["findings"], f"{C_FILES_PER[r][0]:.1f}", CS[r]["strict_hits"], f2(CS[r]["precision_strict"]), f2(CS[r]["precision_lenient"]),
         f2(CS[r]["recall"]), f2(CS[r]["f1_strict"]), pc(CCV[r]["of_223"]), f2(CVAL[r]["adjusted_precision_valid"])] for r in CRUNS]
cur.table(["Model", "Findings", "Per file", "Strict hits", "Precision (strict)", "Precision (lenient)", f"Recall (of {C_N_CONS})", "F1 (strict)", f"Coverage (of {C_N_ALL})", "Adjusted precision"], rows,
          [2.2, 1, 0.9, 1, 1.2, 1.2, 1.1, 1, 1.2, 1.2])
for t_ in T["gen"][1:]:
    cur.para(t_)

cur.h2(T["disc_h"])
for t_ in T["disc"]:
    cur.para(t_)
cur.h2(T["lim_h"])
cur.para(T["lim"])
cur.pagebreak()
first = True
for head, paras_ in T["c6"].items():
    if first:
        cur.h1("CHAPTER 6: CONCLUSION")
        first = False
    cur.h2(head)
    for t_ in paras_:
        cur.para(t_)
cur.pagebreak()

# ---------- 6. References ----------
ref_head = find("References")
ref_els = []
el = ref_head._p.getnext()
while el is not None:
    txt = "".join(t.text or "" for t in el.iter(qn("w:t")))
    if el.tag.endswith("}p") and list(el.iter(qn("w:br"))) and not txt:
        break
    if txt.strip():
        ref_els.append(el)
    el = el.getnext()
old = [P(e).text for e in ref_els]
keep = [re.sub(r"^HashiCorp\. \(2023\)\..*$", "HashiCorp. (2023). HashiCorp state of cloud strategy survey 2023. https://www.hashicorp.com/state-of-the-cloud/2023", t.replace("Anthropic. (2026, February 17)", "Anthropic. (2026b, February 17)")) for t in old if not t.startswith(("Google. (n.d.)", "OpenAI. (n.d.)", "Rahman, A., Parnin, C., & Williams, L. (2023)", "Verdet, M.", "Munshi, A.", "Bridgecrew."))]
bc = [t for t in old if t.startswith("Bridgecrew.")]
for t in bc:
    keep.append(t.replace("(2020)", "(2020a)") if "Checkov" in t else t.replace("(2020)", "(2020b)"))
NEW = [
    "Center for Internet Security. (n.d.). CIS Amazon Web Services Foundations Benchmark. https://www.cisecurity.org/benchmark/amazon_web_services",
    "Chiari, M., De Pascalis, M., & Pradella, M. (2022). Static analysis of infrastructure as code: A survey. In 2022 IEEE 19th International Conference on Software Architecture Companion (ICSA-C) (pp. 218\u2013225). IEEE. https://doi.org/10.1109/ICSA-C54293.2022.00049",
    "Joint Task Force. (2020). Security and privacy controls for information systems and organizations (NIST Special Publication 800-53, Revision 5). National Institute of Standards and Technology. https://doi.org/10.6028/NIST.SP.800-53r5",
    "Amazon Artificial General Intelligence. (2025). The Amazon Nova family of models: Technical report and model card. arXiv. https://arxiv.org/abs/2506.12103",
    "Amazon Web Services. (2026a). Amazon Bedrock user guide. https://docs.aws.amazon.com/bedrock/latest/userguide/",
    "Amazon Web Services. (2026b). Amazon Bedrock pricing. https://aws.amazon.com/bedrock/pricing/",
    "Cohen, J. (1968). Weighted kappa: Nominal scale agreement with provision for scaled disagreement or partial credit. Psychological Bulletin, 70(4), 213–220. https://doi.org/10.1037/h0026256",
    "Hevner, A. R., March, S. T., Park, J., & Ram, S. (2004). Design science in information systems research. MIS Quarterly, 28(1), 75–105. https://doi.org/10.2307/25148625",
    "Kon, P. T. J., Liu, J., Qiu, Y., Fan, W., He, T., Lin, L., Zhang, H., Park, O. M., Elengikal, G. S., Kang, Y., Chen, A., Chowdhury, M., Lee, M., & Wang, X. (2024). IaC-Eval: A code generation benchmark for cloud infrastructure-as-code programs. Advances in Neural Information Processing Systems, 37. https://proceedings.neurips.cc/paper_files/paper/2024/hash/f26b29298ae8acd94bd7e839688e329b-Abstract.html",
    "Rahman, A., Mahdavi-Hezaveh, R., & Williams, L. (2019). A systematic mapping study of infrastructure as code research. Information and Software Technology, 108, 65\u201377. https://doi.org/10.1016/j.infsof.2018.12.004",
    "Runeson, P., & Höst, M. (2009). Guidelines for conducting and reporting case study research in software engineering. Empirical Software Engineering, 14(2), 131–164. https://doi.org/10.1007/s10664-008-9102-8",
    "Saavedra, N., & Ferreira, J. F. (2022). GLITCH: Automated polyglot security smell detection in infrastructure as code. In Proceedings of the 37th IEEE/ACM International Conference on Automated Software Engineering. ACM. https://doi.org/10.1145/3551349.3556945",
    "Sallou, J., Durieux, T., & Panichella, A. (2024). Breaking the silence: The threats of using LLMs in software engineering. In Proceedings of the 46th International Conference on Software Engineering: New Ideas and Emerging Results (ICSE-NIER 2024). ACM. https://arxiv.org/abs/2312.08055",
    "OpenAI. (2026). GPT-5.4 [Large language model]. https://openai.com/",
    "Anthropic. (2026a). Claude Sonnet 5.5 [Large language model]. https://www.anthropic.com/claude",
    "Sokolova, M., & Lapalme, G. (2009). A systematic analysis of performance measures for classification tasks. Information Processing & Management, 45(4), 427–437. https://doi.org/10.1016/j.ipm.2009.03.002",
    "Wohlin, C., Runeson, P., Höst, M., Ohlsson, M. C., Regnell, B., & Wesslén, A. (2012). Experimentation in software engineering. Springer. https://doi.org/10.1007/978-3-642-29044-2",
    "Hou, X., Zhao, Y., Liu, Y., Yang, Z., Wang, K., Li, L., Luo, X., Lo, D., Grundy, J., & Wang, H. (2024). Large language models for software engineering: A systematic literature review. arXiv. https://arxiv.org/abs/2308.10620",
    "Landis, J. R., & Koch, G. G. (1977). The measurement of observer agreement for categorical data. Biometrics, 33(1), 159–174. https://doi.org/10.2307/2529310",
    "Munshi, S., Pathak, S., Ghatode, S., Priyadarshini, T., Chandramouleeswaran, D., & Rana, A. (2025). ACSE-Eval: Can LLMs threat model real-world cloud infrastructure? arXiv. https://arxiv.org/abs/2505.11565",
    "Pearce, H., Ahmad, B., Tan, B., Dolan-Gavitt, B., & Karri, R. (2022). Asleep at the keyboard? Assessing the security of GitHub Copilot's code contributions. In 2022 IEEE Symposium on Security and Privacy (SP). IEEE. https://arxiv.org/abs/2108.09293",
    "Powers, D. M. W. (2011). Evaluation: From precision, recall and F-measure to ROC, informedness, markedness and correlation. Journal of Machine Learning Technologies, 2(1), 37–63.",
    "Rahman, A., Parnin, C., & Williams, L. (2019). The seven sins: Security smells in infrastructure as code scripts. In 2019 IEEE/ACM 41st International Conference on Software Engineering (ICSE) (pp. 164–175). IEEE. https://doi.org/10.1109/ICSE.2019.00033",
    "Ullah, S., Han, M., Pujar, S., Pearce, H., Coskun, A., & Stringhini, G. (2024). LLMs cannot reliably identify and reason about security vulnerabilities (yet?): A comprehensive evaluation, framework, and benchmarks. In 2024 IEEE Symposium on Security and Privacy (SP) (pp. 862–880). IEEE. https://doi.org/10.1109/SP54263.2024.00210",
    "Verdet, A., Hamdaqa, M., Da Silva, L., & Khomh, F. (2025). Assessing the adoption of security policies by developers in terraform across different cloud providers. Empirical Software Engineering, 30, Article 74. https://doi.org/10.1007/s10664-024-10610-0",
    "Zheng, L., Chiang, W.-L., Sheng, Y., Zhuang, S., Wu, Z., Zhuang, Y., Lin, Z., Li, Z., Li, D., Xing, E. P., Zhang, H., Gonzalez, J. E., & Stoica, I. (2023). Judging LLM-as-a-judge with MT-Bench and Chatbot Arena. Advances in Neural Information Processing Systems, 36. https://proceedings.neurips.cc/paper_files/paper/2023/hash/91f18a1287b398d378ef22505bf41832-Abstract.html",
]
allrefs = sorted(keep + NEW, key=lambda t: re.sub(r"[^a-z0-9]", "", t.lower()[:60]))
for i, t in enumerate(allrefs):
    if i < len(ref_els):
        settext(P(ref_els[i]), t)
    else:
        e = copy.deepcopy(ref_els[0]); ref_els[-1].addnext(e); ref_els.append(e); settext(P(e), t)
# put a page break before References

# ---------- 7. Appendices A and B ----------
appA = find("Appendix A: API Cost Estimate")._p
appC = find("Appendix C: Policy ID Mapping")._p
endAB = appC.getprevious()
anchor = appA.getprevious()
delete_between(appA, endAB)
cur = Cur(anchor)
cur.h1("Appendix A: Resource Use and Approximate Cost")
files = sorted((ROOT / "terragoat/terraform/aws").glob("*.tf"))
in_per_model = sum(len(ev.build_prompt(s_, f.read_text(encoding="utf-8"))) for s_ in ("zeroshot", "fewshot") for f in files) // 4
out_tok = {}
for r in RUNS:
    k = r.split("_")[0]
    for f in (ROOT / "responses" / r).glob("*.tf.json"):
        out_tok[k] = out_tok.get(k, 0) + len(json.load(open(f, encoding="utf-8"))["response"]) // 4
n_match = sum(1 for _ in (ROOT / "judgments").glob("*/*.tf.json"))
n_rem = sum(1 for _ in (ROOT / "remediation_scores").glob("*/*/*.json"))
n_val = sum(1 for _ in (ROOT / "validity").glob("*/*.json"))
RATE = {"claude": (3.00, 15.00), "nova": (0.80, 3.20)}
rows_ = []
for k, nm in (("claude", "Claude Sonnet 4.6"), ("gpt", "GPT-5.4"), ("nova", "Amazon Nova Pro")):
    if k in RATE:
        cost = in_per_model * RATE[k][0] / 1e6 + out_tok[k] * RATE[k][1] / 1e6
        rows_.append([nm, f"{in_per_model:,}", f"{out_tok[k]:,}", f"{RATE[k][0]:.2f} / {RATE[k][1]:.2f}", f"{cost:.2f}"])
    else:
        rows_.append([nm, f"{in_per_model:,}", f"{out_tok[k]:,}", "not verified", "not priced"])
cur.para(
    "The evaluation made 84 requests to the three models (14 files, two prompt strategies, three models). The scripts did not record token counts, so Table A.1 "
    "estimates them at four characters per token from the saved prompts and responses. Prices are approximate public list prices for on-demand Amazon Bedrock "
    "use, USD 3.00 and 15.00 per million input and output tokens for Claude Sonnet 4.6 and USD 0.80 and 3.20 for Amazon Nova Pro, taken from third-party price "
    "trackers in October 2026 and not from an invoice. They should be checked against the provider's pricing page (Amazon Web Services, 2026b). A "
    "published rate for GPT-5.4 on Bedrock could not be confirmed, so that model is not priced.")
cur.caption("Table A.1. Estimated tokens and approximate cost of the evaluation requests")
cur.table(["Model", "Input tokens", "Output tokens", "USD per million tokens (input / output)", "Approximate cost (USD)"], rows_, [2.4, 1.4, 1.4, 2.4, 1.8], size=9)
cur.para(
    f"The judging steps added {n_match} matching requests, up to {n_rem} remediation scoring requests and {n_val} validity requests, each carrying one Terraform "
    "file or its candidate list. On these figures the Claude and Nova Pro evaluation requests together cost about "
    f"{sum(float(r_[4]) for r_ in rows_ if r_[4] != 'not priced'):.2f} US dollars, far below the planning estimate of about NZD 100 in the proposal. Actual charges appear in the "
    "billing records of the Amazon Web Services account used.")
cur.pagebreak()
cur.h1("Appendix B: Prompt Templates")
cur.h2("B.1 Zero-shot prompt")
cur.para("The instruction below is followed by the line --- TERRAFORM CONFIGURATION --- and the full text of one file. It is identical for all three models.")
cur.code(ev.build_prompt("zeroshot", "<file content>"))
cur.h2("B.2 Few-shot prompt")
cur.para("The few-shot prompt adds two worked examples between the instruction and the file. The instruction and the file are unchanged.")
cur.code(ev.build_prompt("fewshot", "<file content>"))

# ---------- 8. Appendix D ----------
last = [e for e in body if not e.tag.endswith("sectPr")][-1]
cur = Cur(last)
cur.pagebreak()
cur.h1("Appendix D: Evaluation Records")
cur.para("This appendix documents how the model evaluation was run and scored. Raw responses, judge outputs and score files are stored in the project folder and listed in Section D.8.")
cur.h2("D.1 Run configuration")
cur.table(["Model", "Bedrock identifier", "Temperature", "Max output tokens", "Runs"],
          [["Claude Sonnet 4.6", ev.MODELS["claude"], "0", "16,000", "14 files x 2 prompts"],
           ["GPT-5.4", ev.MODELS["gpt"], "not settable (default)", "16,000", "14 files x 2 prompts"],
           ["Amazon Nova Pro", ev.MODELS["nova"], "0", "9,000", "14 files x 2 prompts"]], [2, 3.2, 1.6, 1.4, 2], size=9)
cur.para(f"All requests were made in the {ev.REGION} region between {FIRST[:16].replace('T', ' ')} and {LAST[:16].replace('T', ' ')}. Each response was saved with its model, strategy, file name and timestamp.")
cur.h2("D.2 Judge prompts")
cur.para("The matching prompt was used with GPT-5.4. The file name, findings and candidate list are inserted where shown.")
cur.code(jf.PROMPT.replace("{{", "{").replace("}}", "}"))
cur.para("The remediation prompt was used with GPT-5.4 and Claude Sonnet 5.5, independently.")
cur.code(rj.PROMPT.replace("{{", "{").replace("}}", "}"))
cur.para("The validity prompt was used with GPT-5.4 on findings that matched no tool finding.")
cur.code(vj.PROMPT.replace("{{", "{").replace("}}", "}"))
cur.h2("D.3 Category assignment rules")
cur.para("Each finding was assigned to the first category whose rule matched, in the order listed. Rules are applied to the check identifier (with the resource name added), the check name and the file name.")
cur.table(["Order", "Category", "Rule (any of)"],
          [["1", "Hardcoded Secrets", "identifier contains SECRET, or name contains HARD and SECRET"],
           ["2", "IAM Over-permissioning", "file name contains IAM, or name contains IAM POLICY, IAM ROLE, PRIVILEGE, CREDENTIALS EXPOSURE, PERMISSION"],
           ["3", "Publicly Exposed Storage", "identifier contains S3 and name contains PUBLIC, BLOCK or ACL"],
           ["4", "Storage Configuration", "identifier or file name contains S3"],
           ["5", "Encryption at Rest", "name contains ENCRYPT, KMS, CMK or ENCRYPTED"],
           ["6", "Network Exposure", "name contains INGRESS, EGRESS, SECURITY GROUP, VPC, SUBNET, PORT 22, PORT 80 or 0.0.0.0"],
           ["7", "Logging & Monitoring", "name contains LOG, LOGGING, MONITORING, CLOUDTRAIL, CLOUDWATCH, AUDIT or TRACING"],
           ["8", "Lambda/Serverless", "file name or name contains LAMBDA"],
           ["9", "Availability & Resilience", "name contains BACKUP, DELETION PROTECTION, MULTI-AZ or MINOR UPGRADE"],
           ["10", "Other", "none of the above"]], [0.6, 2.4, 6], size=9)
cur.h2("D.4 Category-level precision and F1-score")
cur.para("Category precision is the share of a model's findings in that category that matched a consensus finding. A finding with no tool match was placed in a category by applying the rules above to its issue text, so precision here is approximate.")
rows = [[c] + [f"{f2(cat(r, c, 'precision'))} / {f2(cat(r, c, 'f1'))}" for r in RUNS] for c in CATS]
cur.table(["Category (precision / F1)"] + [label(r).replace(", ", "\n") for r in RUNS], rows, [3, 1.2, 1.2, 1.2, 1.2, 1.2, 1.2], size=8)
cur.h2("D.5 Example response and judgement")
ex = json.load(open(ROOT / "responses/claude_zeroshot/s3.tf.json", encoding="utf-8"))
jx = json.load(open(ROOT / "judgments/claude_zeroshot/s3.tf.json", encoding="utf-8"))
first_f = jx["findings"][0]
cur.para("The first finding returned by Claude Sonnet 4.6 for s3.tf under zero-shot prompting is shown with the candidates the judge matched to it.")
cur.code(json.dumps(first_f, indent=1)[:1400])
m0 = next(x for x in jx["matches"] if x["finding"] == 0)
full = {f"c{i}": r for i, r in enumerate(FULL)}
cur.para("Judge output for this finding: " + (", ".join(f"{c} ({full[c]['canonical_check_id']}, {full[c]['status'].lower()})" for c in m0["candidates"]) or "no match") + ".")
cur.h2("D.6 Manual audit of judge decisions")
cur.para("A random sample of 22 matched findings and 30 unmatched findings was read against the Terraform source and the full candidate list. All 22 matches described the same weakness on the same resource as the candidate chosen. "
         "None of the 30 rejected findings had a corresponding candidate for the same resource, including plaintext database credentials in the user data of aws_instance.db_app and the wildcard Elasticsearch "
         "domain policy, for which Checkov and Trivy report no finding. The review was carried out with the help of the AI coding assistant described in Section 3.8, and it should be repeated independently by a second reviewer.")
cur.h2("D.7 Generalisation dataset")
cur.para(f"The second dataset has {C_N_FILES} files ({C_N_GH} from GitHub repositories and {C_N_REG} from Terraform Registry module repositories). Each row gives the repository, the commit that was cloned, the original path and the licence. The files were scanned and evaluated under the identifiers shown.")
cur.table(["Id", "Repository", "Commit", "Original path", "Bytes", "Licence"],
          [[r["id"], r["repository"], r["commit"][:10], r["original_path"], r["bytes"], r["licence"]] for r in CMAN], [0.8, 3.6, 1.2, 2.6, 0.8, 1.2], size=7)
cur.h2("D.8 Files")
cur.table(["File or folder", "Content"],
          [["llm_evaluation.py", "Builds the prompts and calls the models through Amazon Bedrock; saves responses"],
           ["responses/", "Raw model responses, one JSON file per model, strategy and Terraform file (84 files)"],
           ["judge_findings.py, judgments/", "Matching judge and its saved outputs"],
           ["score_results.py", "Computes precision, recall, F1-score and category results"],
           ["scored_summary.csv, scored_by_category.csv, scored_matches.csv", "Scores and the matching decision for every finding"],
           ["extra_stats.py, coverage_summary.csv, paired_tests.csv", "Coverage of all tool findings and exact tests"],
           ["remediation_judge.py, remediation_scores/, aggregate_remediation.py, remediation_summary.csv", "Remediation scoring by two judges and agreement"],
           ["validity_judge.py, validity/, validity_summary.csv", "Exploratory validity check of unmatched findings"],
           ["terragoat/full_tool_comparison.csv, terragoat/consensus_ground_truth.csv", "Checkov and Trivy findings aligned by the crosswalk; the consensus subset"],
           ["collect_corpus.py, scan_corpus.py, corpus_manifest.csv, corpus_eval/", "Collection and scanning of the second dataset and a copy of the evaluation pipeline with its results"],
           ["summarise_validity.py", "Summarises the validity labels"],
           ["bedrock_access_check.py, bedrock_access_check.txt", "Records that each model can be invoked with the settings used"],
           ["make_figures_ch3_app.py", "Draws Figure 3.1 and renders the Appendix E excerpts from saved output"],
           ["build_v05.py, thesis_text.py, thesis_data.py, make_figures_ch5.py", "Scripts that generate Chapter 5 text, figures and this document from the result files"]],
          [4, 6], size=9)

# ---- Appendix E: experiment evidence
cur.pagebreak()
cur.h1("Appendix E: Experiment Evidence")
cur.para("The four figures in this appendix reproduce output from the experiment runs. They are rendered from the saved log and result files named in each caption and "
         "are not screen captures. The full files are in the project folder.")
for fn, cap in [("evidence_E1_run_log.png", "Figure E.1. Excerpt of the saved log from the TerraGoat evaluation run (eval_run.log and eval_nova.log), including the first Nova Pro request error, which was caused by a 16,000 token limit above the model's cap of 10,000 and was fixed by lowering the limit to 9,000."),
                ("evidence_E2_access_check.png", "Figure E.2. Output of bedrock_access_check.py, showing that each of the three models can be invoked through Amazon Bedrock with the settings used in the study (GPT-5.4 does not accept a temperature setting)."),
                ("evidence_E3_scores.png", "Figure E.3. Contents of scored_summary.csv, the scoring output for the six TerraGoat conditions (findings, strict and lenient hits, covered consensus findings, precision, recall and F1-score)."),
                ("evidence_E4_corpus_log.png", "Figure E.4. Excerpt of the saved log from the second-dataset run (corpus_run.log) and the validity summary it ended with.")]:
    cur.figure(ROOT / "terragoat" / fn, cap, 6.2)

# ---- Appendix F: scripts
cur.pagebreak()
cur.h1("Appendix F: Scripts")
cur.para("The scripts below produced the results in Chapter 5. They are listed in the order in which they are run. API credentials are read from an environment variable and are not part of any script.")
SCRIPTS = [("llm_evaluation.py", "Builds the prompts and calls the three models through Amazon Bedrock; saves every response."),
           ("judge_findings.py", "Blinded matching judge: maps each model finding to the tool findings it describes."),
           ("score_results.py", "Computes strict and lenient precision, recall, F1-score and category results."),
           ("extra_stats.py", "Computes coverage of all tool findings and the exact paired tests."),
           ("remediation_judge.py", "Scores remediation advice on the 0 to 2 rubric with two independent judges."),
           ("aggregate_remediation.py", "Aggregates remediation scores and computes judge agreement."),
           ("validity_judge.py", "Exploratory validity labels for findings that matched no tool finding."),
           ("collect_corpus.py", "Collects the GitHub and Registry files for the second dataset."),
           ("scan_corpus.py", "Scans the second dataset with Checkov and Trivy and aligns the findings.")]
for i_, (fn, purpose) in enumerate(SCRIPTS, 1):
    cur.h2(f"F.{i_} {fn}")
    cur.para(purpose)
    lines_ = (ROOT / fn).read_text(encoding="utf-8").splitlines()
    for k_ in range(0, len(lines_), 60):
        pp = cur.code("\n".join(lines_[k_:k_ + 60]))
        for r_ in pp.runs:
            r_.font.size = Pt(7.5)

# ---- Table 4.2: counts from the Chapter 4 category rules (match Figures 4.1 and 4.2)
t42 = doc.tables[2]
for row in t42.rows[1:]:
    fix = {"IAM Over-permissioning": ("12", "5.4%"), "Other": ("55", "24.9%")}.get(row.cells[0].text.strip())
    if fix:
        for cell, val in zip(row.cells[1:3], fix):
            for r_ in cell.paragraphs[0].runs[1:]:
                r_._r.getparent().remove(r_._r)
            cell.paragraphs[0].runs[0].text = val

# ---- remove a citation the source does not support, and disambiguate the two Rahman et al. (2019) works
sub_in(find("Terraform, developed by HashiCorp"), " (Rahman et al., 2019)", "")
for p_ in paras():
    t_ = p_.text
    if "Rahman et al., 2018" in t_:
        sub_in(p_, "Rahman et al., 2018", "Rahman, Mahdavi-Hezaveh, & Williams, 2019")
    if "Rahman et al., 2019" in p_.text:
        sub_in(p_, "Rahman et al., 2019", "Rahman, Parnin, & Williams, 2019")
    if "Rahman et al. (2019)" in p_.text:
        sub_in(p_, "Rahman et al. (2019)", "Rahman, Parnin, and Williams (2019)")

# ---------- 8b. real heading styles so the table of contents lists every chapter ----------
def _styled(p):
    pPr = p._p.pPr
    return pPr is not None and pPr.pStyle is not None
for p in paras():
    t = p.text.strip()
    if not t or len(t) > 100 or _styled(p):
        continue
    lvl = None
    if re.match(r"^Chapter [2-4]: ", t) or t == "References":
        lvl = "Heading 1"
    elif re.match(r"^[2-4]\.\d \S", t) and not t.endswith("."):
        lvl = "Heading 2"
    elif re.match(r"^3\.\d\.\d \S", t) and not t.endswith("."):
        lvl = "Heading 3"
    if lvl == "Heading 1" and t.startswith("Chapter"):
        settext(p, t.upper())
    if lvl:
        p.style = next(st for st in doc.styles if st.style_id == lvl.replace(" ", ""))
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

# ---------- 8c. cover page: programme name and AUT logo ----------
settext(find("Master of Computing and Information Systems"), "Master of Computer and Information Sciences")
first_p = paras()[0]
logo_el = copy.deepcopy(first_p._p)
first_p._p.addprevious(logo_el)
logo_p = P(logo_el)
for r_ in list(logo_p.runs):
    r_._r.getparent().remove(r_._r)
logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
logo_p.paragraph_format.line_spacing = 1.0
logo_p.paragraph_format.space_after = Pt(18)
logo_p.add_run().add_picture(str(ROOT / "aut_logo.png"), width=Inches(1.7))

# ---------- 8d. page breaks: heading-level instead of empty break paragraphs (avoids blank pages) ----------
for el in list(body):
    if not el.tag.endswith("}p"):
        continue
    txt = "".join(t.text or "" for t in el.iter(qn("w:t")))
    if txt.strip() or not list(el.iter(qn("w:br"))) or list(el.iter(qn("w:drawing"))):
        continue
    nxt = el.getnext()
    if nxt is None or not nxt.tag.endswith("}p"):
        continue
    np_ = P(nxt)
    if np_.style is not None and np_.style.name in ("Heading 1",) or np_.text.strip() in ("Abstract", "Table of Contents"):
        el.getparent().remove(el)
        np_.paragraph_format.page_break_before = True

# ---------- 9. table borders, alignment, save ----------
for tb in doc.tables:
    apa_borders(tb)
for p in paras():
    if p.text.strip() and p.style is not None and p.style.name not in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3") and p.paragraph_format.line_spacing == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
doc.core_properties.author = "Prashant Singh"
doc.core_properties.last_modified_by = "Prashant Singh"
doc.core_properties.title = "Evaluating the Effectiveness of Large Language Models for Detecting Security Misconfigurations in Terraform Infrastructure as Code"
doc.save(OUT)
words = sum(len(p.text.split()) for p in doc.paragraphs)
print("saved", OUT.name, "paragraph words:", words)
